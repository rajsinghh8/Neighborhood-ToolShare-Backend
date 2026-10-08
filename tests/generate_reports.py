#!/usr/bin/env python
"""Builds api_test_report.xlsx and project_report.docx from tests-artifacts/test_results.json (observed rows only)."""
import json
import os
from datetime import datetime, timezone

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ART = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tests-artifacts")
rows = json.load(open(os.path.join(ART, "test_results.json")))
total = len(rows)
passed = sum(r["Pass/Fail"] == "PASS" for r in rows)
failed = total - passed
now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
COLS = ["#", "Method", "Endpoint", "Description", "Status Code", "Pass/Fail", "Reason"]

wb = Workbook()
ws = wb.active
ws.title = "Test Results"
ws.append(COLS)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
for r in rows:
    ws.append([r[k] for k in COLS])
    cell = ws.cell(row=ws.max_row, column=6)
    cell.fill = PatternFill("solid", fgColor="C6EFCE" if r["Pass/Fail"] == "PASS" else "FFC7CE")
for col, w in zip("ABCDEFG", (5, 9, 44, 70, 12, 10, 100)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
sm = wb.create_sheet("Summary")
for line in (("Project", "School Management System API"), ("Generated", now), ("Total", total), ("Passed", passed),
             ("Failed", failed), ("New endpoint cases", "GET /api/v1/students/count (rows 48-52)")):
    sm.append(line)
sm.column_dimensions["A"].width = 22
sm.column_dimensions["B"].width = 50
wb.save(os.path.join(ART, "api_test_report.xlsx"))

d = Document()
d.add_heading("School Management System - Project Report", 0)
d.add_paragraph(f"Generated {now}. Spring Boot 3.5.0 / Java 21, Oracle Free 23, Apache Kafka 4.3.1 (KRaft), JWT auth (ADMIN/TEACHER).")
d.add_heading("Result summary", 1)
d.add_paragraph(f"{passed} of {total} checks passed ({failed} failed). All rows are observed results from tests/run_api_tests.py "
                "executed against the running docker compose stack (real Oracle and Kafka).")
d.add_heading("Change in this iteration", 1)
d.add_paragraph("New endpoint GET /api/v1/students/count (ADMIN or TEACHER) returns {\"count\": <long>} via the StudentCountResponse "
                "DTO. Layering: StudentController.count -> StudentService.getStudentCount (read-only transaction, DEBUG log) -> "
                "StudentRepository.count() (inherited from JpaRepository). No schema, config or compose changes.")
d.add_heading("API surface", 1)
for e in ("POST /api/v1/auth/login, /register (ADMIN), /refresh",
          "GET /api/v1/students (paginated offset/limit, grade filter), GET /api/v1/students/count, GET /api/v1/students/{id}",
          "POST/PUT/DELETE /api/v1/students (ADMIN)", "GET /health, /docs, /api-docs (public)",
          "Kafka topic student-events (key = studentId), DLQ student-events.dlq, transactional outbox + idempotent consumer"):
    d.add_paragraph(e, style="List Bullet")
d.add_heading("Test results", 1)
t = d.add_table(rows=1, cols=6)
t.style = "Light Grid Accent 1"
for i, h in enumerate(("#", "Method", "Endpoint", "Description", "Status", "Result")):
    t.rows[0].cells[i].text = h
for r in rows:
    cells = t.add_row().cells
    for i, v in enumerate((r["#"], r["Method"], r["Endpoint"], r["Description"], r["Status Code"], r["Pass/Fail"])):
        cells[i].text = str(v)
failures = [r for r in rows if r["Pass/Fail"] != "PASS"]
d.add_heading("Failures", 1)
d.add_paragraph("None." if not failures else "\n".join(f"#{r['#']} {r['Description']}: {r['Reason']}" for r in failures))
d.save(os.path.join(ART, "project_report.docx"))
print(f"reports written: {passed}/{total} PASS")
